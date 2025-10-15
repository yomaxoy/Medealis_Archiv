"""
Document Generation Service - API Package

Zentralisiertes Document Generation System für DOCX-Dokumente.

Dieses Paket bietet eine vollständige, modulare Lösung für die Dokument-Generierung
mit folgenden Kernkomponenten:

## Hauptmodule:
- document_generation_service: Zentrale API und Service-Schnittstelle
- generation_models: Datenmodelle und Typen für Generation
- document_types: Document Type Definitionen und Template Registry
- context_builder: GenerationContext Management und Erstellung
- word_processor: DOCX-Verarbeitung und Template-Engine
- template_manager: Template-Verwaltung und Caching
- placeholder_engine: Unified Placeholder-System

## Hauptschnittstellen:

### Einfache Document Generation:
```python
from warehouse.application.services.document_generation_service import DocumentGenerationService

service = DocumentGenerationService()
result = service.generate_document(
    document_type="pdb",
    batch_number="P-123456789012",
    delivery_number="DEL-001"
)
```

### Convenience Functions:
```python
from warehouse.application.services.document_generation_service import generate_pdb_document

result = generate_pdb_document(
    batch_number="P-123456789012",
    delivery_number="DEL-001",
    article_number="CT0001"
)
```

### Context-basierte Generation:
```python
from warehouse.application.services.document_generation_service import (
    DocumentGenerationService, create_context_from_parameters, DocumentType
)

context = create_context_from_parameters(
    batch_number="P-123456789012",
    delivery_number="DEL-001",
    document_type=DocumentType.PDB
)

service = DocumentGenerationService()
result = service.generate_document_from_context(context)
```

## Integration mit Phase 1 Storage System:

Das System integriert automatisch mit dem Phase 1 Storage System, falls verfügbar:
- Automatische Pfad-Auflösung
- Storage-Context Integration
- Einheitliche Datei-Verwaltung

## Unterstützte Document Types:
- PDB: Produktdatenblatt
- BEGLEITSCHEIN: Begleitschein
- SICHTKONTROLLE: Sichtkontrolle
- WARENEINGANGSKONTROLLE: Wareneingang
- INCOMING_GOODS_INSPECTION: Incoming Goods Inspection

## Architecture:
Dieses System folgt Clean Architecture Prinzipien mit klarer Trennung von:
- Domain Logic (document_types, generation_models)
- Application Services (document_generation_service, context_builder)
- Infrastructure (template_manager, word_processor)
"""

import logging
from typing import List, Dict, Any, Optional

# =================================================================
# HAUPTSCHNITTSTELLEN - Public API
# =================================================================

# Central Service
from .document_generation_service import (
    DocumentGenerationService,
    generate_pdb_document,
    generate_document_simple,
    validate_generation_service
)

# Core Models und Types
from .generation_models import (
    GenerationContext,
    GenerationResult,
    BatchGenerationResult,
    ProcessingOptions,
    ValidationResult
)

from .document_types import (
    DocumentType,
    TemplateInfo,
    get_template_info,
    get_all_template_info,
    get_supported_document_types,
    TEMPLATE_REGISTRY
)

# Context Management
from .context_builder import (
    ContextBuilder,
    create_context_from_parameters,
    create_context_from_storage,
    validate_context_quality
)

# Processing Engines (für erweiterte Nutzung)
from .word_processor import WordProcessor
from .template_manager import TemplateManager
from .placeholder_engine import PlaceholderEngine

logger = logging.getLogger(__name__)

# =================================================================
# PACKAGE METADATA
# =================================================================

__version__ = "1.0.0"
__author__ = "Medealis System"
__description__ = "Centralized Document Generation Service for DOCX documents"

# =================================================================
# PACKAGE CONFIGURATION
# =================================================================

# Default Service Configuration
DEFAULT_CONFIG = {
    'template_base_dir': None,  # Auto-detect
    'default_output_format': 'docx',
    'auto_open_documents': False,
    'strict_validation': False,
    'enable_caching': True,
    'storage_integration': True
}

# =================================================================
# CONVENIENCE FUNCTIONS - Package Level API
# =================================================================

def create_service(config: Optional[Dict[str, Any]] = None) -> DocumentGenerationService:
    """
    Erstellt DocumentGenerationService mit Konfiguration.

    Args:
        config: Service-Konfiguration (optional)

    Returns:
        Konfigurierte DocumentGenerationService Instanz
    """
    try:
        service = DocumentGenerationService()

        if config:
            service.configure_service(config)

        return service

    except Exception as e:
        logger.error(f"Failed to create DocumentGenerationService: {e}")
        raise


def get_available_document_types() -> List[str]:
    """
    Gibt Liste aller verfügbaren Document Types zurück.

    Returns:
        Liste der Document Type Namen
    """
    try:
        return [doc_type.value for doc_type in DocumentType]
    except Exception as e:
        logger.error(f"Error getting document types: {e}")
        return []


def get_template_information() -> Dict[str, Dict[str, Any]]:
    """
    Gibt detaillierte Template-Informationen zurück.

    Returns:
        Dictionary mit Template-Informationen für jeden Document Type
    """
    try:
        template_info = {}

        for doc_type in DocumentType:
            info = get_template_info(doc_type)
            if info:
                template_info[doc_type.value] = {
                    'name': info.name,
                    'filename': info.filename,
                    'description': info.description,
                    'version': info.version,
                    'required_placeholders': info.required_placeholders,
                    'optional_placeholders': info.optional_placeholders,
                    'supported_formats': info.supported_formats,
                    'has_tables': info.has_tables,
                    'supports_conditional_content': info.supports_conditional_content
                }

        return template_info

    except Exception as e:
        logger.error(f"Error getting template information: {e}")
        return {}


def validate_package_health() -> Dict[str, Any]:
    """
    Validiert Package-Gesundheit und Abhängigkeiten.

    Returns:
        Health Check Ergebnis
    """
    try:
        health_check = {
            'package_healthy': True,
            'version': __version__,
            'dependencies': {},
            'components': {},
            'errors': [],
            'warnings': []
        }

        # Check python-docx availability
        try:
            from docx import Document
            health_check['dependencies']['python_docx'] = {
                'available': True,
                'version': getattr(Document, '__version__', 'unknown')
            }
        except ImportError:
            health_check['dependencies']['python_docx'] = {'available': False}
            health_check['errors'].append("python-docx not available - DOCX processing disabled")
            health_check['package_healthy'] = False

        # Check Storage System Integration
        try:
            from ..document_storage.document_storage_service import DocumentStorageService
            health_check['dependencies']['storage_system'] = {'available': True}
        except ImportError:
            health_check['dependencies']['storage_system'] = {'available': False}
            health_check['warnings'].append("Phase 1 Storage System not available")

        # Component Health Checks
        try:
            service = DocumentGenerationService()
            component_health = service.validate_service_health()
            health_check['components'] = component_health.get('components', {})

            if not component_health.get('service_healthy', False):
                health_check['package_healthy'] = False
                health_check['errors'].extend(component_health.get('errors', []))

        except Exception as e:
            health_check['package_healthy'] = False
            health_check['errors'].append(f"Service initialization failed: {e}")

        return health_check

    except Exception as e:
        return {
            'package_healthy': False,
            'error': str(e),
            'version': __version__
        }


def get_package_statistics() -> Dict[str, Any]:
    """
    Gibt Package-Statistiken zurück.

    Returns:
        Dictionary mit Package-Statistiken
    """
    try:
        stats = {
            'package_version': __version__,
            'supported_document_types': len(DocumentType),
            'available_templates': len(get_all_template_info()),
            'template_registry_size': len(TEMPLATE_REGISTRY)
        }

        # Template Details
        template_details = {}
        for doc_type in DocumentType:
            info = get_template_info(doc_type)
            if info:
                template_details[doc_type.value] = {
                    'required_placeholders': len(info.required_placeholders),
                    'optional_placeholders': len(info.optional_placeholders),
                    'total_placeholders': len(info.get_all_placeholders())
                }

        stats['template_details'] = template_details

        return stats

    except Exception as e:
        logger.error(f"Error getting package statistics: {e}")
        return {'error': str(e)}


# =================================================================
# PACKAGE INITIALIZATION
# =================================================================

def _initialize_package():
    """Initialisiert Package beim Import."""
    try:
        logger.info(f"Document Generation Service Package v{__version__} loaded")

        # Validiere kritische Abhängigkeiten
        try:
            from docx import Document
        except ImportError:
            logger.warning("python-docx not available - Document generation will be limited")

        # Storage System Check
        try:
            from ..document_storage.document_storage_service import DocumentStorageService
            logger.info("Phase 1 Storage System integration available")
        except ImportError:
            logger.info("Phase 1 Storage System not available - using standalone mode")

    except Exception as e:
        logger.error(f"Package initialization failed: {e}")


# Initialize package when imported
_initialize_package()

# =================================================================
# PUBLIC API EXPORTS
# =================================================================

__all__ = [
    # Main Service
    'DocumentGenerationService',
    'create_service',

    # Core Models
    'GenerationContext',
    'GenerationResult',
    'BatchGenerationResult',
    'ProcessingOptions',
    'ValidationResult',

    # Document Types
    'DocumentType',
    'TemplateInfo',
    'get_template_info',
    'get_all_template_info',
    'get_supported_document_types',
    'TEMPLATE_REGISTRY',

    # Context Management
    'ContextBuilder',
    'create_context_from_parameters',
    'create_context_from_storage',
    'validate_context_quality',

    # Processing Components (advanced)
    'WordProcessor',
    'TemplateManager',
    'PlaceholderEngine',

    # Convenience Functions
    'generate_pdb_document',
    'generate_document_simple',

    # Package Management
    'get_available_document_types',
    'get_template_information',
    'validate_package_health',
    'validate_generation_service',
    'get_package_statistics',

    # Configuration
    'DEFAULT_CONFIG'
]

# =================================================================
# DEPRECATION NOTICES
# =================================================================

# Future API Changes
DEPRECATED_FUNCTIONS = {
    # Placeholder for future deprecations
}

def _check_deprecated_usage():
    """Prüft auf veraltete API-Nutzung."""
    # Implementation für Deprecation Warnings
    pass